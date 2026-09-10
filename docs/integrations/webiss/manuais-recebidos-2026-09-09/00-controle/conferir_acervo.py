from pathlib import Path
import hashlib,json,subprocess,zipfile,xml.etree.ElementTree as E
import openpyxl
root=Path(__file__).resolve().parent.parent
rows=[]
for p in sorted(root.rglob('*')):
 if not p.is_file() or p.relative_to(root).parts[0] in ['00-controle','06-textos-extraidos']:continue
 data=p.read_bytes();row={'path':str(p.relative_to(root)),'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()}
 if p.suffix=='.pdf':
  r=subprocess.run(['pdfinfo',str(p)],capture_output=True,text=True,check=True);row['validation']='PDF legível'
 elif p.suffix=='.zip':
  with zipfile.ZipFile(p) as z:assert z.testzip() is None;row['entries']=len(z.infolist());row['validation']='CRC ZIP OK'
 elif p.suffix in ['.xml','.xsd','.wsdl']:
  d=E.fromstring(data);row['xml_root']=d.tag;row['validation']='XML bem formado'
 elif p.suffix=='.xlsx':
  w=openpyxl.load_workbook(p,read_only=True,data_only=False);row['sheets']=[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in w];w.close();row['validation']='XLSX legível'
 elif p.suffix=='.png':row['validation']='Evidência do usuário preservada'
 rows.append(row)
(root/'00-controle/inventario-integridade.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
(root/'00-controle/SHA256SUMS.txt').write_text('\n'.join(r['sha256']+'  '+r['path'] for r in rows)+'\n')
print('Arquivos de acervo conferidos:',len(rows))
for r in rows:
 if 'sheets' in r:print(r['sheets'])
